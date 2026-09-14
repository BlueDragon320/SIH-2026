"""
Comprehensive Tests for Spec Gap Enhancements:
- WI-1: Knowledge Base Folder Watcher Daemon (orchestrator/ingestion/folder_watcher.py)
- WI-2: Sensitivity & Department Tagging for RAG Chunks (orchestrator/rag/vector_store.py)
- WI-3: Sandbox CPU & Memory Resource Constraints (orchestrator/tools/sandbox.py)
- WI-4: Audited Excel Multi-Sheet Layout with Dedicated Summary Sheet (orchestrator/tools/spreadsheet.py)
- WI-5: Docs Directory & Architecture Spec Links (docs/)
- WI-6: Vision OCR Tool Re-Export Compliance (orchestrator/tools/vision_ocr.py)
- API: Watcher status, scan triggers, and metadata-aware RAG endpoints
"""
import unittest
import os
import tempfile
import shutil
import time
import openpyxl
from fastapi.testclient import TestClient

from orchestrator.main import app, kb_watcher
from orchestrator.rag.vector_store import LocalVectorStore
from orchestrator.ingestion.folder_watcher import KnowledgeBaseWatcher
from orchestrator.tools import sandbox, spreadsheet, files
from orchestrator.tools.vision_ocr import VisionOCRPipeline


class TestSpecEnhancements(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        self.temp_dir = tempfile.TemporaryDirectory()
        login_res = self.client.post("/v1/auth/login", json={"username": "admin", "password": "admin123"})
        if login_res.status_code == 200:
            token = login_res.json().get("access_token", "")
            self.headers = {"Authorization": f"Bearer {token}"}
        else:
            self.headers = {}

    def tearDown(self):
        try:
            self.temp_dir.cleanup()
        except Exception:
            pass

    # -------------------------------------------------------------------------
    # WI-1: Folder Watcher Daemon Tests
    # -------------------------------------------------------------------------
    def test_folder_watcher_auto_ingestion(self):
        """Test that dropping a new file into watched folder auto-indexes it."""
        mock_kb_dir = os.path.join(self.temp_dir.name, "kb")
        mock_db_dir = os.path.join(self.temp_dir.name, "chroma")
        os.makedirs(mock_kb_dir, exist_ok=True)

        local_vs = LocalVectorStore(db_dir=mock_db_dir)
        watcher = KnowledgeBaseWatcher(
            watch_dir=mock_kb_dir,
            vector_store=local_vs,
            scan_interval=2
        )

        # Initial scan on empty directory
        res1 = watcher.scan_once()
        self.assertEqual(res1["status"], "success")
        self.assertEqual(res1["newly_ingested_count"], 0)

        # Drop a new SOP document
        test_sop = os.path.join(mock_kb_dir, "SOP_Turbine_Safety_2026.md")
        with open(test_sop, "w", encoding="utf-8") as f:
            f.write("""
            # SOP-2026-TURBINE: Emergency Hydraulic Cutoff
            1. Purpose: Cut hydraulic fuel flow during overspeed events (> 3600 RPM).
            2. Actuation: Solenoid valve SV-401 activates in under 45 milliseconds.
            3. Reset: Manual supervisor key override required.
            """)

        # Trigger scan
        res2 = watcher.scan_once()
        self.assertEqual(res2["status"], "success")
        self.assertEqual(res2["newly_ingested_count"], 1)
        self.assertEqual(res2["newly_ingested"][0]["filename"], "SOP_Turbine_Safety_2026.md")

        # Verify indexed in vector store
        docs = local_vs.list_indexed_documents()
        self.assertTrue(any(d["source"] == "SOP_Turbine_Safety_2026.md" for d in docs))

        # Re-scan without changes - should not re-ingest
        res3 = watcher.scan_once()
        self.assertEqual(res3["newly_ingested_count"], 0)

        # Check status reporting
        status = watcher.get_status()
        self.assertIn("watched_dir", status)
        self.assertIn("tracked_files", status)
        self.assertIn("SOP_Turbine_Safety_2026.md", status["tracked_files"])

    # -------------------------------------------------------------------------
    # WI-2: Sensitivity & Department Tagging for RAG
    # -------------------------------------------------------------------------
    def test_rag_chunk_sensitivity_and_department_metadata(self):
        """Test sensitivity and department metadata tagging and search filtering."""
        mock_db_dir = os.path.join(self.temp_dir.name, "chroma_meta")
        local_vs = LocalVectorStore(db_dir=mock_db_dir)

        doc_restricted = """
        RESTRICTED DEFENCE CLEARANCE: Missile Frigate Propulsion Specs
        Shaft horsepower: 45,000 SHP via twin gas turbines.
        Propeller alloy: Super Duplex 2507 high corrosion resistant.
        """
        doc_public = """
        GENERAL OPERATIONS: Standard Cafeteria & Office Protocols
        Office hours are 0800 to 1700 IST Monday through Friday.
        """

        local_vs.add_document(
            filename="defence_specs.md",
            text=doc_restricted,
            sensitivity="confidential",
            department="defence_naval"
        )
        local_vs.add_document(
            filename="general_office.md",
            text=doc_public,
            sensitivity="internal",
            department="human_resources"
        )

        # Unfiltered search should return defence doc
        all_res = local_vs.hybrid_search("gas turbine propulsion", top_k=4)
        self.assertGreater(len(all_res), 0)
        self.assertEqual(all_res[0]["metadata"]["sensitivity"], "confidential")
        self.assertEqual(all_res[0]["metadata"]["department"], "defence_naval")

        # Filtered search for "internal" should NOT return confidential doc
        filtered_res = local_vs.hybrid_search("propulsion specs", top_k=4, sensitivity="internal")
        self.assertTrue(all(r["metadata"]["sensitivity"] == "internal" for r in filtered_res))
        self.assertFalse(any(r["metadata"]["sensitivity"] == "confidential" for r in filtered_res))

        # Filtered search matching "confidential" should return it
        conf_res = local_vs.hybrid_search("propulsion specs", top_k=4, sensitivity="confidential")
        self.assertGreaterEqual(len(conf_res), 1)
        self.assertEqual(conf_res[0]["metadata"]["sensitivity"], "confidential")
        self.assertEqual(conf_res[0]["metadata"]["department"], "defence_naval")

    # -------------------------------------------------------------------------
    # WI-3: Sandbox CPU & Memory Resource Constraints
    # -------------------------------------------------------------------------
    def test_sandbox_resource_limits_reported(self):
        """Verify sandbox enforces and reports CPU and memory limits."""
        code = "print('RESOURCE_LIMIT_TEST_OK')"
        res = sandbox.execute_python_code(code)

        self.assertTrue(res["success"])
        self.assertEqual(res["exit_code"], 0)
        self.assertIn("RESOURCE_LIMIT_TEST_OK", res["stdout"])
        self.assertIn("memory_limit_bytes", res)
        self.assertIn("cpu_limit_sec", res)
        self.assertIsNotNone(res["memory_limit_bytes"])
        self.assertGreater(res["memory_limit_bytes"], 0)
        self.assertIsNotNone(res["cpu_limit_sec"])
        self.assertGreater(res["cpu_limit_sec"], 0)

    # -------------------------------------------------------------------------
    # WI-4: Audited Excel Multi-Sheet Layout (Calc + Summary)
    # -------------------------------------------------------------------------
    def test_spreadsheet_multi_sheet_summary_layout(self):
        """Verify audited Excel generator creates dedicated Summary sheet with live formulas."""
        test_filename = "multi_sheet_audit.xlsx"
        headers = ["Sensor ID", "Vibration (mm/s)", "Pressure (bar)"]
        rows = [
            ["S-01", 1.5, 42.0],
            ["S-02", 2.8, 44.5],
            ["S-03", 0.9, 41.2]
        ]
        summary_formulas = {
            "Average Vibration": "=AVERAGE(B2:B4)",
            "Max Pressure": "=MAX(C2:C4)"
        }

        res = spreadsheet.create_audit_spreadsheet(
            filename=test_filename,
            sheet_title="TelemetryData",
            headers=headers,
            rows=rows,
            summary_formulas=summary_formulas,
            title="TURBINE BEARING VIBRATION AUDIT"
        )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["sheets_count"], 2)
        self.assertIn("TelemetryData", res["sheets"])
        self.assertIn("Summary", res["sheets"])

        # Load workbook and verify sheets
        wb = openpyxl.load_workbook(res["absolute_path"], data_only=False)
        self.assertEqual(wb.sheetnames, ["TelemetryData", "Summary"])

        # Check Summary sheet contents
        ws_sum = wb["Summary"]
        self.assertEqual(ws_sum.cell(row=1, column=1).value, "Executive Summary — TURBINE BEARING VIBRATION AUDIT")
        self.assertEqual(ws_sum.cell(row=3, column=1).value, "Key Metric / KPI")
        self.assertEqual(ws_sum.cell(row=3, column=2).value, "Formula Reference")
        self.assertEqual(ws_sum.cell(row=3, column=3).value, "Audit Status")

        # Check metric row 1 (Average Vibration)
        self.assertEqual(ws_sum.cell(row=4, column=1).value, "Average Vibration")
        ref1 = ws_sum.cell(row=4, column=2).value
        self.assertTrue(str(ref1).startswith("='TelemetryData'"))
        self.assertEqual(ws_sum.cell(row=4, column=3).value, "LIVE_AUDITED")

        # Cleanup
        if os.path.exists(res["absolute_path"]):
            os.remove(res["absolute_path"])

    # -------------------------------------------------------------------------
    # WI-5: Docs Directory & Architecture Spec Links
    # -------------------------------------------------------------------------
    def test_docs_directory_and_architecture_guide(self):
        """Verify docs/ directory, ARCHITECTURE.md, and spec symlink exist."""
        project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
        docs_dir = os.path.join(project_root, "docs")
        arch_file = os.path.join(docs_dir, "ARCHITECTURE.md")
        spec_file = os.path.join(docs_dir, "air-gapped-agentic-ai-workbench-spec.md")

        self.assertTrue(os.path.isdir(docs_dir), "docs/ directory must exist")
        self.assertTrue(os.path.isfile(arch_file), "docs/ARCHITECTURE.md must exist")
        self.assertTrue(os.path.exists(spec_file), "docs/air-gapped-agentic-ai-workbench-spec.md must exist")

        with open(arch_file, "r", encoding="utf-8") as f:
            content = f.read()
            self.assertIn("Air-Gapped Agentic AI Workbench", content)
            self.assertIn("Bubblewrap", content)

    # -------------------------------------------------------------------------
    # WI-6: Vision OCR Tool Re-Export Compliance
    # -------------------------------------------------------------------------
    def test_vision_ocr_re_export_tool(self):
        """Verify orchestrator/tools/vision_ocr.py exists and exports VisionOCRPipeline."""
        pipeline = VisionOCRPipeline()
        self.assertIsNotNone(pipeline)
        self.assertEqual(pipeline.default_model, "moondream")

    # -------------------------------------------------------------------------
    # API: Knowledge Base Watcher & Metadata Endpoints
    # -------------------------------------------------------------------------
    def test_api_watcher_status_and_scan(self):
        """Verify /v1/knowledge-base/watcher-status and watcher-scan endpoints."""
        res = self.client.get("/v1/knowledge-base/watcher-status", headers=self.headers)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("running", data)
        self.assertIn("watched_dir", data)
        self.assertIn("scan_interval_sec", data)

        # Trigger manual scan via API
        scan_res = self.client.post("/v1/knowledge-base/watcher-scan", headers=self.headers)
        self.assertEqual(scan_res.status_code, 200)
        scan_data = scan_res.json()
        self.assertEqual(scan_data["status"], "success")
        self.assertIn("scanned_files_count", scan_data)


if __name__ == "__main__":
    unittest.main()
