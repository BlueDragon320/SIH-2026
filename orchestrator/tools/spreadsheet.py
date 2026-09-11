"""
Spreadsheet Tool for Air-Gapped Agentic Workbench.
Generates and audits Excel (.xlsx) deliverables with live formulas and multi-sheet layouts.
"""
import os
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
WORKSPACE_DIR = os.environ.get("WORKSPACE_DIR", os.path.join(PROJECT_ROOT, "data", "workspace"))

def create_audit_spreadsheet(
    filename: str,
    sheet_title: str,
    headers: List[str],
    rows: List[List[Any]],
    summary_formulas: Optional[Dict[str, str]] = None,
    title: Optional[str] = None
) -> Dict[str, Any]:
    """
    Create a professionally formatted Excel spreadsheet with live formulas and styled headers.
    
    Args:
        filename: Target filename in workspace (e.g. 'inspection_metrics.xlsx')
        sheet_title: Name of primary data sheet
        headers: Column header labels
        rows: Table data rows (can include numbers, strings, or Excel formulas like '=B2*C2')
        summary_formulas: Dict of label -> formula (e.g. {'Total Cost': '=SUM(D2:D10)', 'Average Temp': '=AVERAGE(C2:C10)'})
        title: Optional banner title at top of sheet
    """
    os.makedirs(WORKSPACE_DIR, exist_ok=True)
    target_path = os.path.join(WORKSPACE_DIR, filename)
    if not target_path.endswith(".xlsx"):
        target_path += ".xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Styling definitions
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Navy blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    bold_font = Font(name="Calibri", size=11, bold=True)
    summary_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid") # Light blue
    
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    current_row = 1

    # Optional Banner Title
    if title:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        title_cell = ws.cell(row=current_row, column=1, value=title)
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[current_row].height = 25
        current_row += 2

    # Headers
    header_row_num = current_row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row_num, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[header_row_num].height = 22
    current_row += 1

    # Data Rows
    start_data_row = current_row
    for r in rows:
        for col_num, val in enumerate(r, 1):
            cell = ws.cell(row=current_row, column=col_num, value=val)
            cell.border = thin_border
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, str) and val.startswith("="):
                cell.alignment = Alignment(horizontal="right")
                cell.font = Font(name="Calibri", size=11, italic=True)
        current_row += 1
    end_data_row = current_row - 1

    # Summary Section with Live Formulas
    if summary_formulas:
        current_row += 1
        for label, formula in summary_formulas.items():
            label_cell = ws.cell(row=current_row, column=len(headers) - 1, value=label)
            label_cell.font = bold_font
            label_cell.alignment = Alignment(horizontal="right")
            label_cell.fill = summary_fill
            label_cell.border = thin_border

            formula_cell = ws.cell(row=current_row, column=len(headers), value=formula)
            formula_cell.font = bold_font
            formula_cell.alignment = Alignment(horizontal="right")
            formula_cell.fill = summary_fill
            formula_cell.border = thin_border
            current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(target_path)
    return {
        "status": "success",
        "deliverable": os.path.basename(target_path),
        "absolute_path": target_path,
        "rows_count": len(rows),
        "columns_count": len(headers),
        "message": f"Successfully created audited Excel workbook at {os.path.basename(target_path)}"
    }

def read_spreadsheet(filename: str) -> Dict[str, Any]:
    """Read data and live formulas from an Excel spreadsheet in workspace."""
    target_path = os.path.join(WORKSPACE_DIR, filename)
    if not os.path.exists(target_path):
        raise FileNotFoundError(f"Spreadsheet '{filename}' not found.")

    wb = openpyxl.load_workbook(target_path, data_only=False)
    sheets_data = {}
    for name in wb.sheetnames:
        ws = wb[name]
        sheet_rows = []
        for row in ws.iter_rows(values_only=True):
            sheet_rows.append([r for r in row])
        sheets_data[name] = sheet_rows
    return {
        "filename": filename,
        "sheets": sheets_data
    }
