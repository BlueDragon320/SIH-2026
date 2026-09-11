"""
Document Generation Tool for Air-Gapped Agentic Workbench.
Creates professional Word (.docx) approval notes/reports and PowerPoint (.pptx) decks.
"""
import os
import datetime
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor

WORKSPACE_DIR = os.path.abspath("/home/blue/SIH/data/workspace")

def _set_cell_background(cell, fill_hex: str):
    """Set background color of a docx table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_approval_note(
    filename: str,
    title: str,
    reference_no: Optional[str] = None,
    department: str = "Engineering & Operations",
    author: str = "Lead Inspection Officer",
    background: str = "",
    findings: List[str] = None,
    findings_table: Optional[Dict[str, Any]] = None,
    risk_assessment: str = "",
    recommendations: List[str] = None,
    signoff_name: str = "Chief Technical Advisor"
) -> Dict[str, Any]:
    """
    Generate an authentic enterprise/defence-grade Approval Note (.docx).
    """
    os.makedirs(WORKSPACE_DIR, exist_ok=True)
    target_path = os.path.join(WORKSPACE_DIR, filename)
    if not target_path.endswith(".docx"):
        target_path += ".docx"

    findings = findings or []
    recommendations = recommendations or []
    ref_no = reference_no or f"REF/ENG/{datetime.datetime.now().strftime('%Y%m%d')}/01"
    today_str = datetime.date.today().strftime("%d %B %Y")

    doc = Document()

    # Document Header Banner
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ref_run = header_para.add_run(f"REF: {ref_no}\nDATE: {today_str}\nCONFIDENTIALITY: RESTRICTED / AIR-GAPPED INTERNAL")
    ref_run.font.size = Pt(8.5)
    ref_run.font.color.rgb = RGBColor(120, 120, 120)

    # Document Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("OFFICIAL APPROVAL NOTE")
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 73, 125)

    sub_title_para = doc.add_paragraph()
    sub_title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_title_para.add_run(title.upper())
    sub_run.font.size = Pt(13)
    sub_run.font.bold = True

    # Metadata Info Block
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [f"Originating Department: {department}", f"Prepared By: {author}"],
        [f"Reference Index: {ref_no}", f"Target Action: Immediate Approval"]
    ]
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = meta_data[r_idx][c_idx]
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            _set_cell_background(cell, "F2F2F2")

    doc.add_paragraph() # Spacing

    # Section 1: Background
    h1 = doc.add_heading("1. Background & Context", level=1)
    h1.style.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph(background or "This note summarizes recent inspections, test observations, and compliance assessments.")

    # Section 2: Key Findings
    h2 = doc.add_heading("2. Technical Observations & Findings", level=1)
    h2.style.font.color.rgb = RGBColor(31, 73, 125)
    for f in findings:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(f)
        r.font.size = Pt(10.5)

    # Optional Structured Table in Findings
    if findings_table and "headers" in findings_table and "rows" in findings_table:
        headers = findings_table["headers"]
        rows = findings_table["rows"]
        t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col_idx, h_text in enumerate(headers):
            cell = t.cell(0, col_idx)
            cell.text = str(h_text)
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _set_cell_background(cell, "1F497D")
        for row_idx, r_data in enumerate(rows):
            for col_idx, val in enumerate(r_data):
                cell = t.cell(row_idx + 1, col_idx)
                cell.text = str(val)
                cell.paragraphs[0].runs[0].font.size = Pt(9.5)
                if row_idx % 2 == 1:
                    _set_cell_background(cell, "F9FAFB")
        doc.add_paragraph() # Spacing

    # Section 3: Risk Assessment
    if risk_assessment:
        h3 = doc.add_heading("3. Risk & Impact Assessment", level=1)
        h3.style.font.color.rgb = RGBColor(31, 73, 125)
        doc.add_paragraph(risk_assessment)

    # Section 4: Recommendations
    h4 = doc.add_heading("4. Proposed Recommendations & Next Steps", level=1)
    h4.style.font.color.rgb = RGBColor(31, 73, 125)
    for rec in recommendations:
        p = doc.add_paragraph(style='List Number')
        r = p.add_run(rec)
        r.font.size = Pt(10.5)

    # Section 5: Sign-off & Authorizations Block
    doc.add_paragraph()
    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.cell(0, 0).text = "Submitted By:\n\n_______________________\n" + author
    sign_table.cell(0, 1).text = "Approved By:\n\n_______________________\n" + signoff_name
    sign_table.cell(1, 0).text = f"Date: {today_str}"
    sign_table.cell(1, 1).text = f"Date: {today_str} (Air-Gap Digital Seal)"

    doc.save(target_path)
    return {
        "status": "success",
        "deliverable": os.path.basename(target_path),
        "absolute_path": target_path,
        "type": "docx",
        "message": f"Generated official Approval Note: {os.path.basename(target_path)}"
    }

def create_presentation(
    filename: str,
    title: str,
    subtitle: str = "Air-Gapped Engineering Summary",
    slides: Optional[List[Dict[str, Any]]] = None
) -> Dict[str, Any]:
    """
    Generate an executive PowerPoint (.pptx) deck.
    """
    os.makedirs(WORKSPACE_DIR, exist_ok=True)
    target_path = os.path.join(WORKSPACE_DIR, filename)
    if not target_path.endswith(".pptx"):
        target_path += ".pptx"

    slides = slides or []
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"{subtitle}\nGenerated On-Premises | Air-Gapped"

    # Bullet Slides
    bullet_slide_layout = prs.slide_layouts[1]
    for s_info in slides:
        s_title = s_info.get("title", "Slide")
        bullets = s_info.get("bullets", [])
        slide = prs.slides.add_slide(bullet_slide_layout)
        slide.shapes.title.text = s_title
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for idx, bullet in enumerate(bullets):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
            p.text = bullet
            p.font.size = PptxPt(18)
            p.level = 0

    prs.save(target_path)
    return {
        "status": "success",
        "deliverable": os.path.basename(target_path),
        "absolute_path": target_path,
        "type": "pptx",
        "slides_count": len(slides) + 1,
        "message": f"Generated PowerPoint Presentation: {os.path.basename(target_path)}"
    }
